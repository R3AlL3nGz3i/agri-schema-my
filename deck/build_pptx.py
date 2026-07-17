#!/usr/bin/env python3
"""Render the Open Design deck (deck/index.html) slide-by-slide to PNGs and
assemble a real 16:9 .pptx (one full-bleed image per slide).

Reproducible: re-run after Open Design iterates on index.html.

    .venv/bin/python deck/build_pptx.py

Deps (already in .venv): playwright (+chromium), python-pptx.
"""
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

DECK_DIR = Path(__file__).resolve().parent
HTML = DECK_DIR / "index.html"
SLIDES_DIR = DECK_DIR / "slides"
PPTX_OUT = DECK_DIR / "agrischema-hackathon.pptx"

# 1920x1080 base, rendered at 2x for crisp text in the .pptx.
BASE_W, BASE_H = 1920, 1080
SCALE = 2

# EMU per inch = 914400. 16:9 canvas.
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5


def render_slides() -> list[Path]:
    SLIDES_DIR.mkdir(exist_ok=True)
    paths: list[Path] = []
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(
            viewport={"width": BASE_W, "height": BASE_H},
            device_scale_factor=SCALE,
        )
        page.goto(HTML.as_uri())
        # Wait for webfonts so text metrics are final.
        page.wait_for_load_state("networkidle")
        page.evaluate("document.fonts && document.fonts.ready")

        # Freeze the stage at 1:1 (drop the fit() scale) and hide nav chrome.
        page.evaluate(
            """() => {
                const stage = document.getElementById('deck-stage');
                stage.style.transform = 'none';
                stage.style.boxShadow = 'none';
                for (const sel of ['.deck-counter', '.deck-hint']) {
                    const el = document.querySelector(sel);
                    if (el) el.style.display = 'none';
                }
            }"""
        )

        count = page.evaluate("document.querySelectorAll('.slide').length")
        stage = page.locator("#deck-stage")
        for i in range(count):
            page.evaluate(
                """(i) => {
                    const slides = document.querySelectorAll('.slide');
                    slides.forEach((el, j) => el.classList.toggle('active', j === i));
                }""",
                i,
            )
            page.evaluate("document.fonts && document.fonts.ready")
            page.wait_for_timeout(200)
            out = SLIDES_DIR / f"slide-{i + 1:02d}.png"
            stage.screenshot(path=str(out))
            paths.append(out)
            print(f"  rendered {out.name}")
        browser.close()
    return paths


def build_pptx(images: list[Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    blank = prs.slide_layouts[6]  # fully blank layout
    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(img), 0, 0, width=prs.slide_width, height=prs.slide_height
        )
    prs.save(PPTX_OUT)
    print(f"wrote {PPTX_OUT} ({len(images)} slides)")


if __name__ == "__main__":
    print(f"Rendering slides from {HTML.name} ...")
    imgs = render_slides()
    build_pptx(imgs)
